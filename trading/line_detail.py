"""Line Detail enrichment for trading (BRIEF #2).

Attaches Plank's canonical product attributes -- department / item_type /
style / collection / finish / lifecycle status / kit / GM -- to trading
order+SKU lines, sourced from the Line Detail (the authoritative product
reference; see reference/line_detail_data_dictionary_v3.pdf and
plank_domain_glossary_v2.pdf). This is a pure read-and-join layer: it does
not compute revenue (trading/revenue.py's AB formula) and never drops a
line from the reconciled total -- unmatched SKUs stay in, tagged
enriched=False; filters like is_live_* / is_el_component are FLAGS a
consumer applies, not row-drops (BRIEF #2 §6).

Deliberately NOT built on common/sku_taxonomy.py: that module derives
`department` from the SKU's own code (first-letter lookup), with Line
Detail only informing item_type/style, and a lenient sku-code fallback
when Line Detail has no row -- the right behaviour for returns, where the
glossary explicitly wants a fallback. This brief wants the opposite for
trading: department comes ONLY from Line Detail's own resolved "Product
Type" field (see §7: "Don't derive department from the code"), and an
unmatched SKU is tagged "Unknown", not code-guessed. Two different,
deliberate answers to "what's this SKU's department" for two different
reports -- noted here so it isn't mistaken for drift.

LINE_DETAIL_SOURCE switches where the workbook comes from; both paths
converge on the same load_line_detail(path) parser -- "dropbox" (later)
just needs to land a file shaped like the committed local snapshot
(header on row 1, same column names) before calling it, so no parsing
logic ever branches on source.
"""
import os
import sys
from datetime import datetime, date

import openpyxl

LINE_DETAIL_SOURCE = os.environ.get("LINE_DETAIL_SOURCE", "local")  # "local" | "dropbox"
LINE_DETAIL_LOCAL_PATH = os.path.join(os.path.dirname(__file__), "source", "line_detail.xlsx")

# §4 lifecycle vocabulary -- pinned. Any non-blank value outside this set
# aborts (build_line_detail_index raises ValueError naming the SKU + value).
# Blank/None status is treated as "unspecified", not a vocabulary violation --
# a handful of real Line Detail rows are incomplete admin entries, not typos;
# see build_line_detail_index's docstring for the reasoning.
STATUS_ENUM = {
    "In Development", "Launching", "Live", "Not For Sale",
    "Discontinued", "Dead", "Disco to Resource", "Not Sold in this Market",
}

ADS_CLASS_ENUM = {"CORE", "STANDARD", "EXCLUDE"}
MARKET_SCOPE_ENUM = {"US", "GB", "Global"}

NEWNESS_MONTHS = 6

# canonical field -> source column header, per BRIEF #2 §3. Product
# Type/Category/Sub Category are read straight off these headers even
# though (per the glossary's reversed-label gotcha) "Product Type" holds
# department-grain values ("Cabinetry") and "Product Category" holds
# item_type-grain values ("Handle") -- i.e. read literally, not by what the
# header text sounds like it should mean.
COLUMN_MAP = {
    "sku": "SKU",
    "description": "Product Description",
    "status_uk": "UK Status",
    "status_us": "US Status",
    "is_kit_raw": "Kit?",
    "is_assembly_raw": "Assembled SKU?",
    "department": "Product Type",
    "item_type": "Product Category",
    "style": "Sub Category",
    "collection": "Collection",
    "material": "Material",
    "finish": "Finish",
    "market_scope": "US/GB/Global SKU",
    "supplier_cost_gbp": "Supplier Cost Price (£)",
    "rrp_gbp_incvat": "RRP in GBP (£)",
    "rrp_usd_extax": "RRP in USD ($)",
    "launch_date_raw": "Launch Date",
    "ads_class": "Included or Excluded from Ads",
}

UNKNOWN = "Unknown"


def _s(v):
    """Normalise any cell value to a stripped string, '' for blank/None."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _bool_yes_no(v, sku, field_name, warnings):
    s = _s(v)
    if s == "":
        return False
    if s.lower() == "yes":
        return True
    if s.lower() == "no":
        return False
    warnings.append(f"{sku}: unexpected {field_name} value {v!r}, treating as False")
    return False


def parse_launch_date(raw):
    """Best-effort parse of the free-text Launch Date field (§4). Returns a
    `date` or None if unparseable. Real data is mostly already a datetime
    (Excel-native dates survive openpyxl's data_only load); the one
    confirmed free-text pattern is literal "pre-2022"-style strings, which
    fall through to None here -- the caller defaults unparseable dates to
    Continuity and counts them, per §4's explicit instruction not to
    silently drop them.
    """
    if raw is None or raw == "":
        return None
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    s = str(raw).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m", "%B %Y", "%b %Y", "%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _months_before(d, months):
    """d minus `months` calendar months (day-of-month clamped to 28 to
    dodge month-length edge cases -- newness is a 6-month bucket, not a
    billing date, so day-level precision doesn't matter).
    """
    year = d.year
    month = d.month - months
    while month <= 0:
        month += 12
        year -= 1
    return date(year, month, min(d.day, 28))


def load_line_detail(path=None):
    """Load the committed/fetched Line Detail workbook -> list of raw dict
    records keyed by the canonical field names in COLUMN_MAP, in file order
    (not yet de-duped or validated -- build_line_detail_index does that).
    """
    path = path or LINE_DETAIL_LOCAL_PATH
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["Line Detail"]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    col_idx = {name: header.index(src) for name, src in COLUMN_MAP.items()}

    records = []
    for row in rows:
        sku = _s(row[col_idx["sku"]])
        if not sku:
            continue
        record = {name: row[col_idx[name]] for name in COLUMN_MAP}
        record["sku"] = sku
        records.append(record)
    return records


def build_line_detail_index(path=None, as_of=None):
    """De-dupe Line Detail to one row per SKU, validate the status enum,
    and derive the per-SKU flags (is_live_*, newness_bucket, gm_pct) ->
    {sku: enriched_record}.

    as_of: the date newness_bucket is computed relative to (typically the
    report month's end, e.g. date(2026, 5, 31) for a May 2026 build) --
    NOT wall-clock "today", so a monthly report's Newness/Continuity split
    doesn't drift every time it's re-run later. Required if any record is
    live in either market.

    De-dupe rule: first occurrence wins (matches common/enrichment.py's
    line_detail_lookup -- pandas drop_duplicates default). The only
    duplicate SKUs found in the real file are literal "--NAME-0-" template
    placeholders for still-in-development products with no SKU assigned
    yet (e.g. "--KEPL-0-" reused across several draft rows) -- these never
    appear on a real order line, so which draft row wins is immaterial.
    """
    records = load_line_detail(path)
    warnings = []

    by_sku = {}
    dup_count = 0
    for r in records:
        if r["sku"] in by_sku:
            dup_count += 1
            continue
        by_sku[r["sku"]] = r
    if dup_count:
        print(f"line_detail: de-duped {dup_count} duplicate-SKU row(s) (first occurrence kept)",
              file=sys.stderr)

    blank_status_count = 0
    unparsed_launch_count = 0
    index = {}

    for sku, r in by_sku.items():
        status_uk = _s(r["status_uk"])
        status_us = _s(r["status_us"])
        for field, value in (("UK Status", status_uk), ("US Status", status_us)):
            if value == "":
                blank_status_count += 1
            elif value not in STATUS_ENUM:
                raise ValueError(
                    f"line_detail: {sku} has {field}={value!r}, outside the §4 enum {sorted(STATUS_ENUM)}"
                )

        is_live_uk = status_uk == "Live"
        is_live_us = status_us == "Live"

        launch_date = parse_launch_date(r["launch_date_raw"])
        newness_bucket = None
        if is_live_uk or is_live_us:
            if launch_date is None:
                unparsed_launch_count += 1
                newness_bucket = "Continuity"
            else:
                if as_of is None:
                    raise ValueError("build_line_detail_index: as_of is required when any SKU is live")
                newness_bucket = "Newness" if launch_date >= _months_before(as_of, NEWNESS_MONTHS) else "Continuity"

        market_scope = _s(r["market_scope"]) or None
        if market_scope is not None and market_scope not in MARKET_SCOPE_ENUM:
            warnings.append(f"{sku}: unexpected market_scope value {market_scope!r}")

        ads_class = _s(r["ads_class"]) or None
        if ads_class is not None and ads_class not in ADS_CLASS_ENUM:
            warnings.append(f"{sku}: unexpected ads_class value {ads_class!r}")

        rrp_gbp_incvat = r["rrp_gbp_incvat"]
        rrp_exvat = (rrp_gbp_incvat / 1.20) if isinstance(rrp_gbp_incvat, (int, float)) else None
        supplier_cost_gbp = r["supplier_cost_gbp"] if isinstance(r["supplier_cost_gbp"], (int, float)) else None
        gm_pct = None
        if rrp_exvat and supplier_cost_gbp is not None:
            gm_pct = (rrp_exvat - supplier_cost_gbp) / rrp_exvat

        index[sku] = {
            "sku": sku,
            "description": _s(r["description"]) or None,
            "department": _s(r["department"]) or UNKNOWN,
            "item_type": _s(r["item_type"]) or UNKNOWN,
            "style": _s(r["style"]) or UNKNOWN,
            "collection": _s(r["collection"]) or UNKNOWN,
            "material": _s(r["material"]) or None,
            "finish": _s(r["finish"]) or None,
            "status_uk": status_uk or None,
            "status_us": status_us or None,
            "is_live_uk": is_live_uk,
            "is_live_us": is_live_us,
            "newness_bucket": newness_bucket,
            "market_scope": market_scope,
            "is_kit": _bool_yes_no(r["is_kit_raw"], sku, "Kit?", warnings),
            "is_assembly": _bool_yes_no(r["is_assembly_raw"], sku, "Assembled SKU?", warnings),
            "is_el_component": _s(r["collection"]).upper() == "EL COMPONENT",
            "ads_class": ads_class,
            "supplier_cost_gbp": supplier_cost_gbp,
            "rrp_gbp_incvat": rrp_gbp_incvat if isinstance(rrp_gbp_incvat, (int, float)) else None,
            "rrp_usd_extax": r["rrp_usd_extax"] if isinstance(r["rrp_usd_extax"], (int, float)) else None,
            "rrp_exvat": rrp_exvat,
            "gm_pct": gm_pct,
            "enriched": True,
        }

    if blank_status_count:
        print(f"line_detail: {blank_status_count} UK/US Status field(s) blank across "
              f"{len(index)} SKUs (treated as unspecified, not an enum violation)", file=sys.stderr)
    if unparsed_launch_count:
        print(f"line_detail: {unparsed_launch_count} live SKU(s) had an unparseable Launch Date "
              f"(e.g. 'pre-2022') -- defaulted to Continuity per §4", file=sys.stderr)
    for w in warnings:
        print(f"line_detail: WARNING {w}", file=sys.stderr)

    return index


_UNMATCHED_RECORD = {
    "description": None, "department": UNKNOWN, "item_type": UNKNOWN, "style": UNKNOWN,
    "collection": UNKNOWN, "material": None, "finish": None, "status_uk": None, "status_us": None,
    "is_live_uk": False, "is_live_us": False, "newness_bucket": None, "market_scope": None,
    "is_kit": False, "is_assembly": False, "is_el_component": False, "ads_class": None,
    "supplier_cost_gbp": None, "rrp_gbp_incvat": None, "rrp_usd_extax": None, "rrp_exvat": None,
    "gm_pct": None, "enriched": False,
}


def enrich_lines(lines, index):
    """Left-join trading lines (matrixify_source.build_lines() output) onto
    the Line Detail index on `sku`, normalised (stripped). Never drops a
    line -- an unmatched SKU keeps the line with `enriched=False` and
    department/collection/item_type/style all "Unknown" (§5), so the
    reconciled revenue total is untouched by enrichment.
    """
    enriched = []
    for line in lines:
        sku = _s(line.get("sku"))
        attrs = index.get(sku)
        if attrs is None:
            attrs = dict(_UNMATCHED_RECORD)
            attrs["sku"] = sku or None
        enriched.append({**line, **{k: v for k, v in attrs.items() if k != "sku"}})
    return enriched


def coverage_report(enriched_lines, ab_of_line):
    """Unmatched-SKU diagnostic (§10 acceptance check #1): count and total
    revenue of lines that didn't match Line Detail. `ab_of_line(line)` is
    the caller's AB-per-line function (kept out of this module -- it does
    not compute revenue), so this can be reused regardless of which FX
    rate/formula produced the line's revenue contribution.
    """
    total_ab = matched_ab = 0.0
    total_lines = matched_lines = 0
    unmatched_skus = {}
    for line in enriched_lines:
        ab = ab_of_line(line)
        total_ab += ab
        total_lines += 1
        if line["enriched"]:
            matched_ab += ab
            matched_lines += 1
        else:
            unmatched_skus[line.get("sku")] = unmatched_skus.get(line.get("sku"), 0.0) + ab

    coverage_pct = (matched_ab / total_ab * 100) if total_ab else 100.0
    return {
        "total_lines": total_lines,
        "matched_lines": matched_lines,
        "unmatched_lines": total_lines - matched_lines,
        "total_ab": total_ab,
        "matched_ab": matched_ab,
        "unmatched_ab": total_ab - matched_ab,
        "coverage_pct": coverage_pct,
        "unmatched_skus": unmatched_skus,
    }


if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else None
    as_of = date(2026, 5, 31)
    index = build_line_detail_index(path, as_of=as_of)
    print(f"\nLine Detail index: {len(index)} distinct SKUs")
    depts = {}
    for rec in index.values():
        depts[rec["department"]] = depts.get(rec["department"], 0) + 1
    print("department counts:", depts)
